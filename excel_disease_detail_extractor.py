"""Extract western diagnostic details from the May 5 disease Excel.

The source sheet mixes western diagnosis notes with TCM syndrome, formulas,
herbs, dosing, and pharmacy plan text. This extractor keeps only M1-useful
western diagnosis evidence and writes review-only cache files. It never writes
to diseases_core.json.
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import disease_manifest_builder as manifest_builder


BASE_DIR = Path(__file__).resolve().parent
CACHE_DIR = BASE_DIR / "data" / "disease_cache"
DEFAULT_EXCEL = manifest_builder.DEFAULT_EXCEL
DETAILS_PATH = CACHE_DIR / "excel_disease_diagnostic_details.json"
REPORT_PATH = CACHE_DIR / "excel_disease_detail_report.md"
PENDING_PATH = CACHE_DIR / "pending_disease_cards.json"

FORBIDDEN_DETAIL_TERMS = [
    "方剂", "基础方", "中药", "草药", "治法", "治则", "辨证", "证型", "病机",
    "处方", "用药", "剂量", "加减", "服法", "煎服", "药店", "终端", "颗粒",
    "汤剂", "丸", "散", "膏", "贴", "针灸", "推拿", "穴位", "经络",
    "舌象", "脉象", "舌红", "舌淡", "舌暗", "苔黄", "苔白", "滑脉", "数脉",
    "气虚", "阴虚", "阳虚", "血虚", "本虚", "标实", "寒证", "热证", "痰热",
    "伏痰", "脏腑", "扶正", "祛邪", "外邪", "风寒", "风热", "湿热", "内饮",
    "寒饮", "痰饮", "肺脾", "肺肾", "脾肾", "肝肾", "肾阳", "肾阴", "虚火",
]

WESTERN_KEEP_TERMS = [
    "西医", "诊断", "鉴别", "并发症", "临床表现", "症状", "体征", "检查",
    "实验室", "影像", "指标", "血常规", "尿常规", "CRP", "PCT", "FeNO", "EOS",
    "IgE", "CT", "MRI", "X线", "胸片", "超声", "心电图", "肺功能", "PEF",
    "SpO2", "血氧", "发热", "咳嗽", "喘息", "胸闷", "鼻塞", "流涕", "腹痛",
    "腹泻", "呕吐", "皮疹", "出血", "水肿", "疼痛", "感染", "炎症", "肿大",
    "阻塞", "狭窄", "损伤", "风险", "危险", "急性", "慢性",
]

CHECK_TERMS = [
    "血常规", "尿常规", "便常规", "CRP", "PCT", "FeNO", "EOS", "IgE",
    "影像", "胸片", "X线", "CT", "MRI", "超声", "心电图", "肺功能", "PEF",
    "SpO2", "血氧", "病原", "培养", "抗原", "抗体", "核酸", "支原体",
]

RED_FLAG_TERMS = [
    "低氧", "呼吸困难", "发绀", "休克", "意识", "抽搐", "脱水", "高热",
    "持续发热", "咯血", "胸痛", "喘憋", "SpO2", "并发症", "危重", "急重",
]

AXIS_TERMS = {
    "infection_or_inflammation": ["感染", "炎症", "发热", "CRP", "PCT", "白细胞"],
    "allergic_th2_inflammation": ["过敏", "变应", "嗜酸", "EOS", "IgE", "FeNO"],
    "upper_airway_involvement": ["鼻", "咽", "喉", "扁桃体", "腺样体", "上气道"],
    "lower_airway_involvement": ["肺", "支气管", "喘息", "啰音", "胸片", "下气道"],
    "airway_hyperreactivity": ["气道高反应", "喘息", "胸闷", "PEF", "肺功能"],
    "secretion_or_retention": ["分泌", "痰", "鼻涕", "流涕", "黏液"],
    "mucosal_barrier_dysfunction": ["黏膜", "糜烂", "溃疡", "屏障"],
    "structural_abnormality": ["肥大", "狭窄", "阻塞", "畸形", "结构"],
    "autoimmune_or_immune_complex": ["自身免疫", "免疫复合物", "IgA", "抗体"],
    "coagulation_or_bleeding_disorder": ["出血", "凝血", "紫癜", "血小板"],
    "urinary_renal_risk": ["肾", "尿", "蛋白尿", "血尿"],
    "digestive_reflux_or_dysfunction": ["胃", "肠", "腹泻", "呕吐", "反流"],
    "cardiovascular_or_circulatory_risk": ["心", "血压", "循环", "胸痛"],
    "neoplastic_or_proliferative": ["肿瘤", "癌", "增生", "占位"],
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("**", "").replace("###", "").replace("##", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _split_detail_units(text: str) -> list[str]:
    text = _clean(text)
    if not text:
        return []
    text = re.sub(r"([。；;])", r"\1\n", text)
    text = re.sub(r"(【[^】]{1,30}】)", r"\n\1", text)
    text = re.sub(r"(\d+[.、])", r"\n\1", text)
    units = []
    for raw in text.splitlines():
        item = raw.strip(" \t-•|")
        if item:
            units.append(item)
    return units


def _is_forbidden_detail(text: str) -> bool:
    if any(term in text for term in FORBIDDEN_DETAIL_TERMS):
        return True
    if re.search(r"\b\d+(\.\d+)?\s*(g|mg|ml|片|粒|袋|丸)\b", text, re.I):
        return True
    return False


def _is_western_detail(text: str) -> bool:
    if _is_forbidden_detail(text):
        return False
    return any(term in text for term in WESTERN_KEEP_TERMS)


def _add_unique(target: list[str], value: str, max_len: int = 220) -> None:
    value = _clean(value)
    value = re.sub(r"^[-—:：\d.、 ]+", "", value)
    if not value or len(value) > max_len:
        return
    if _is_forbidden_detail(value):
        return
    if value not in target:
        target.append(value)


def _classify_units(units: list[str]) -> dict[str, list[str]]:
    result = {
        "diagnostic_key_points": [],
        "clinical_manifestations": [],
        "differential_diagnoses_text": [],
        "complications": [],
        "required_checks": [],
        "red_flags": [],
        "objective_markers": [],
        "raw_western_detail_excerpt": [],
    }
    for unit in units:
        if not _is_western_detail(unit):
            continue
        _add_unique(result["raw_western_detail_excerpt"], unit, max_len=320)
        if "鉴别" in unit:
            _add_unique(result["differential_diagnoses_text"], unit)
        if "并发症" in unit or "并发" in unit:
            _add_unique(result["complications"], unit)
        if any(term in unit for term in CHECK_TERMS):
            _add_unique(result["required_checks"], unit)
            _add_unique(result["objective_markers"], unit)
        if any(term in unit for term in RED_FLAG_TERMS):
            _add_unique(result["red_flags"], unit)
        if any(term in unit for term in ["症状", "体征", "临床表现", "关键体征", "发热", "咳嗽", "喘息", "疼痛", "皮疹", "呕吐", "腹泻", "鼻塞", "流涕"]):
            _add_unique(result["clinical_manifestations"], unit)
        if any(term in unit for term in ["诊断", "西医", "分期", "指标", "检查", "影像", "实验室"]):
            _add_unique(result["diagnostic_key_points"], unit)

    # If Excel only provides manifestations/checks, they still help rule-in.
    for field in ("clinical_manifestations", "required_checks", "objective_markers"):
        for item in result[field][:4]:
            _add_unique(result["diagnostic_key_points"], item)
    return result


def _infer_axes(detail: dict[str, list[str]], name: str) -> list[dict[str, Any]]:
    blob = name + " " + json.dumps(detail, ensure_ascii=False)
    axes: list[dict[str, Any]] = []
    for axis, terms in AXIS_TERMS.items():
        hits = [term for term in terms if term in blob]
        if hits:
            axes.append({
                "axis": axis,
                "role": "primary" if not axes else "secondary",
                "weight": 0.75 if not axes else 0.45,
                "confidence": "medium" if len(hits) >= 2 else "low",
                "typical_findings": hits[:5],
            })
    return axes[:6]


def _headers(row: tuple[Any, ...]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for i, cell in enumerate(row):
        value = _clean(cell.value)
        if value:
            headers[value] = i
    return headers


def _empty_detail_entry(standard: str, raw_name: str, row, header: dict[str, int], row_number: int) -> dict[str, Any]:
    def field(name: str) -> str:
        idx = header.get(name)
        return _clean(row[idx].value) if idx is not None and idx < len(row) else ""

    return {
        "standard_western_name": standard,
        "aliases": [],
        "source_raw_names": [raw_name] if raw_name else [],
        "department_or_category": field("科目"),
        "sex": field("性别"),
        "typical_age": field("好发年龄"),
        "entry_type": manifest_builder._entry_type(standard),
        "m1_diagnosis_entry_allowed": manifest_builder._allowed(standard),
        "diagnostic_key_points": [],
        "clinical_manifestations": [],
        "differential_diagnoses_text": [],
        "complications": [],
        "required_checks": [],
        "red_flags": [],
        "objective_markers": [],
        "pathology_axes": [],
        "raw_western_detail_excerpt": [],
        "source_row_numbers": [row_number],
        "source_file": str(manifest_builder.resolve_excel_path(DEFAULT_EXCEL)),
        "status": "pending_review",
        "confidence": "low",
    }


def _merge_list(entry: dict[str, Any], field: str, values: list[Any]) -> None:
    current = entry.setdefault(field, [])
    for value in values:
        if value and value not in current:
            current.append(value)


def build_excel_diagnostic_details(excel_path: str | Path = DEFAULT_EXCEL) -> dict[str, Any]:
    try:
        path = manifest_builder.resolve_excel_path(excel_path)
    except FileNotFoundError:
        if Path(excel_path) == DEFAULT_EXCEL and DETAILS_PATH.exists():
            return json.loads(DETAILS_PATH.read_text(encoding="utf-8"))
        raise
    manifest = manifest_builder.build_disease_manifest(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    details_by_name: dict[str, dict[str, Any]] = {}
    raw_detail_cells = 0
    rows_with_kept_detail = 0

    for ws in wb.worksheets:
        rows = ws.iter_rows()
        try:
            first = next(rows)
        except StopIteration:
            continue
        header = _headers(first)
        disease_idx = header.get("病名【诊断】")
        detail_idx = header.get("全病程管理方案（备份待删除）")
        if disease_idx is None:
            continue

        for offset, row in enumerate(rows, start=2):
            disease_cell = row[disease_idx].value if disease_idx < len(row) else None
            if not disease_cell:
                continue
            names, aliases, _warnings = manifest_builder._extract_names(_clean(disease_cell))
            raw_detail = row[detail_idx].value if detail_idx is not None and detail_idx < len(row) else None
            raw_detail_text = _clean(raw_detail)
            if raw_detail_text:
                raw_detail_cells += 1
            classified = _classify_units(_split_detail_units(raw_detail_text))
            if classified["raw_western_detail_excerpt"]:
                rows_with_kept_detail += 1

            for raw_name in names:
                standard, normalization_aliases = manifest_builder._standardize(raw_name)
                if not standard:
                    continue
                entry = details_by_name.get(standard)
                if entry is None:
                    entry = _empty_detail_entry(standard, raw_name, row, header, offset)
                    details_by_name[standard] = entry
                elif offset not in entry["source_row_numbers"]:
                    entry["source_row_numbers"].append(offset)
                _merge_list(entry, "aliases", [a for a in aliases + normalization_aliases if a and a != standard])
                _merge_list(entry, "source_raw_names", [raw_name])
                for field, values in classified.items():
                    _merge_list(entry, field, values)

    manifest_names = {item["standard_western_name"]: item for item in manifest["diseases"]}
    for standard, item in manifest_names.items():
        if standard not in details_by_name:
            details_by_name[standard] = {
                "standard_western_name": standard,
                "aliases": item.get("aliases", []),
                "source_raw_names": [item.get("raw_text", standard)],
                "department_or_category": item.get("department_or_category", ""),
                "sex": item.get("sex", ""),
                "typical_age": item.get("typical_age", ""),
                "entry_type": item.get("entry_type", "standard_western_disease"),
                "m1_diagnosis_entry_allowed": item.get("m1_diagnosis_entry_allowed", True),
                "diagnostic_key_points": [],
                "clinical_manifestations": [],
                "differential_diagnoses_text": [],
                "complications": [],
                "required_checks": [],
                "red_flags": [],
                "objective_markers": [],
                "pathology_axes": [],
                "raw_western_detail_excerpt": [],
                "source_row_numbers": [],
                "source_file": str(path),
                "status": "pending_review",
                "confidence": "low",
            }

    for entry in details_by_name.values():
        entry["pathology_axes"] = _infer_axes(entry, entry["standard_western_name"])
        useful_count = sum(len(entry.get(field, [])) for field in (
            "diagnostic_key_points", "clinical_manifestations", "required_checks",
            "objective_markers", "differential_diagnoses_text", "complications",
        ))
        entry["confidence"] = "medium" if useful_count >= 4 and entry["pathology_axes"] else "low"

    diseases = sorted(details_by_name.values(), key=lambda item: item["standard_western_name"])
    result = {
        "created_at": _now(),
        "source_file": str(path),
        "raw_row_count": manifest["raw_row_count"],
        "raw_disease_cell_count": manifest["raw_disease_cell_count"],
        "cleaned_disease_count": len(diseases),
        "raw_detail_cell_count": raw_detail_cells,
        "rows_with_kept_western_detail": rows_with_kept_detail,
        "diseases_with_any_western_detail": sum(1 for d in diseases if d["raw_western_detail_excerpt"]),
        "diseases_with_required_checks": sum(1 for d in diseases if d["required_checks"]),
        "diseases_with_differential_text": sum(1 for d in diseases if d["differential_diagnoses_text"]),
        "diseases_with_complications": sum(1 for d in diseases if d["complications"]),
        "diseases": diseases,
    }
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    DETAILS_PATH.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    _merge_details_into_pending_cards(diseases)
    _write_report(result)
    return result


def _pending_card_from_detail(detail: dict[str, Any]) -> dict[str, Any]:
    checks = detail.get("required_checks", [])
    manifestations = detail.get("clinical_manifestations", [])
    diagnostic = detail.get("diagnostic_key_points", [])
    rule_in = (manifestations + detail.get("objective_markers", []) + diagnostic)[:8]
    rule_out = detail.get("differential_diagnoses_text", [])[:4]
    if not rule_out:
        rule_out = ["Excel未提供明确排除项，需结合鉴别诊断矩阵和客观检查复核"]
    return {
        "disease_name": detail["standard_western_name"],
        "aliases": detail.get("aliases", []),
        "source_raw_names": detail.get("source_raw_names", []),
        "department_or_category": detail.get("department_or_category", ""),
        "entry_type": detail.get("entry_type", "standard_western_disease"),
        "m1_diagnosis_entry_allowed": detail.get("m1_diagnosis_entry_allowed", True),
        "diagnostic_key_points": diagnostic[:8],
        "rule_in_features": rule_in[:8],
        "rule_out_features": rule_out[:6],
        "differential_diagnoses": [
            {"disease_name": "待人工结构化鉴别诊断", "distinguishing_points": detail.get("differential_diagnoses_text", [])[:4]}
        ] if detail.get("differential_diagnoses_text") else [],
        "required_checks": checks[:8],
        "red_flags": detail.get("red_flags", [])[:8],
        "typical_age_or_population": [v for v in [detail.get("typical_age", ""), detail.get("sex", "")] if v],
        "typical_course": "",
        "pathology_axes": detail.get("pathology_axes", [])[:6],
        "must_not_be_primary_when": [],
        "need_external_check_if": ["Excel诊断字段不足或鉴别诊断未结构化"] if detail.get("confidence") == "low" else [],
        "source_summary": ["5月5日知识库Excel：仅抽取西医诊断、表现、检查、鉴别和并发症相关字段；已过滤药物方案与传统医学内容"],
        "source_urls": [],
        "confidence": detail.get("confidence", "low"),
        "status": "pending_review",
        "created_at": _now(),
        "updated_at": _now(),
        "excel_diagnostic_detail_source": {
            "source_file": detail.get("source_file", ""),
            "source_row_numbers": detail.get("source_row_numbers", []),
        },
    }


def _load_pending() -> list[dict[str, Any]]:
    if not PENDING_PATH.exists():
        return []
    try:
        data = json.loads(PENDING_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except json.JSONDecodeError:
        return []


def _sanitize_pending_value(value: Any) -> Any:
    if isinstance(value, str):
        return "" if any(term in value for term in FORBIDDEN_DETAIL_TERMS) else value
    if isinstance(value, list):
        cleaned = [_sanitize_pending_value(item) for item in value]
        return [item for item in cleaned if item not in ("", [], {})]
    if isinstance(value, dict):
        cleaned_dict = {key: _sanitize_pending_value(item) for key, item in value.items()}
        return {key: item for key, item in cleaned_dict.items() if item not in ("", [], {})}
    return value


def _ensure_pending_schema(card: dict[str, Any]) -> dict[str, Any]:
    defaults = {
        "disease_name": "",
        "aliases": [],
        "source_raw_names": [],
        "department_or_category": "",
        "entry_type": "standard_western_disease",
        "m1_diagnosis_entry_allowed": True,
        "diagnostic_key_points": [],
        "rule_in_features": [],
        "rule_out_features": [],
        "differential_diagnoses": [],
        "required_checks": [],
        "red_flags": [],
        "typical_age_or_population": [],
        "typical_course": "",
        "pathology_axes": [],
        "must_not_be_primary_when": [],
        "need_external_check_if": [],
        "source_summary": [],
        "source_urls": [],
        "confidence": "low",
        "status": "pending_review",
        "created_at": _now(),
        "updated_at": _now(),
    }
    for key, value in defaults.items():
        card.setdefault(key, value)
    return card


def _merge_details_into_pending_cards(details: list[dict[str, Any]]) -> None:
    valid_detail_names = {detail["standard_western_name"] for detail in details}
    cards = [
        _ensure_pending_schema(card)
        for card in (_sanitize_pending_value(card) for card in _load_pending())
        if (
            isinstance(card, dict)
            and card.get("disease_name")
            and (
                not card.get("excel_diagnostic_detail_source")
                or card.get("disease_name") in valid_detail_names
            )
        )
    ]
    by_name = {card.get("disease_name"): card for card in cards if card.get("disease_name")}
    for detail in details:
        name = detail["standard_western_name"]
        if not detail.get("m1_diagnosis_entry_allowed", True):
            continue
        incoming = _pending_card_from_detail(detail)
        existing = by_name.get(name)
        if existing is None:
            cards.append(incoming)
            by_name[name] = incoming
            continue
        for field in (
            "aliases", "source_raw_names", "diagnostic_key_points", "rule_in_features",
            "rule_out_features", "required_checks", "red_flags", "typical_age_or_population",
            "pathology_axes", "need_external_check_if", "source_summary",
        ):
            _merge_list(existing, field, incoming.get(field, []))
        if incoming.get("differential_diagnoses"):
            existing.setdefault("differential_diagnoses", [])
            for item in incoming["differential_diagnoses"]:
                if item not in existing["differential_diagnoses"]:
                    existing["differential_diagnoses"].append(item)
        existing["pathology_axes"] = existing.get("pathology_axes", [])[:6]
        existing["status"] = "pending_review"
        existing["updated_at"] = _now()
        existing["excel_diagnostic_detail_source"] = incoming["excel_diagnostic_detail_source"]
    PENDING_PATH.write_text(json.dumps(cards, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_report(result: dict[str, Any]) -> None:
    sparse = [
        d["standard_western_name"]
        for d in result["diseases"]
        if not d.get("raw_western_detail_excerpt")
    ][:80]
    lines = [
        "# Excel Disease Diagnostic Detail Report",
        "",
        f"- generated_at: {result['created_at']}",
        f"- source_file: {result['source_file']}",
        f"- raw_disease_cell_count: {result['raw_disease_cell_count']}",
        f"- cleaned_disease_count: {result['cleaned_disease_count']}",
        f"- raw_detail_cell_count: {result['raw_detail_cell_count']}",
        f"- rows_with_kept_western_detail: {result['rows_with_kept_western_detail']}",
        f"- diseases_with_any_western_detail: {result['diseases_with_any_western_detail']}",
        f"- diseases_with_required_checks: {result['diseases_with_required_checks']}",
        f"- diseases_with_differential_text: {result['diseases_with_differential_text']}",
        f"- diseases_with_complications: {result['diseases_with_complications']}",
        "",
        "## Scope",
        "",
        "Only western diagnosis-useful content was extracted: diagnosis, clinical manifestations, objective checks, differential text, complications, red flags, and coarse pathology axes.",
        "Medication, formulas, herbs, doses, prescriptions, pharmacy plans, TCM syndromes, TCM pathogenesis, tongue, and pulse text were filtered out.",
        "",
        "## Sparse Entries Needing Manual Or External Review",
        "",
    ]
    lines.extend(f"- {name}" for name in sparse)
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    output = build_excel_diagnostic_details()
    print(json.dumps({
        "cleaned_disease_count": output["cleaned_disease_count"],
        "diseases_with_any_western_detail": output["diseases_with_any_western_detail"],
        "diseases_with_required_checks": output["diseases_with_required_checks"],
        "diseases_with_differential_text": output["diseases_with_differential_text"],
        "diseases_with_complications": output["diseases_with_complications"],
    }, ensure_ascii=False, indent=2))
