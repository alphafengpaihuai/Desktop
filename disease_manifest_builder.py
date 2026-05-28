"""Build disease_manifest.json from the May 5 Excel knowledge base.

Only western diagnosis metadata is extracted. Formula, herbs, dose, treatment
plans, and TCM syndrome text are ignored by design.
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
CACHE_DIR = BASE_DIR / "data" / "disease_cache"
MANIFEST_PATH = CACHE_DIR / "disease_manifest.json"
DEFAULT_EXCEL = Path("/Users/fangxuan/Desktop/5月5日 知识库 病名lsx.xlsx")

ENTRY_TYPES = {
    "症状": "symptom_or_sign",
    "体征": "symptom_or_sign",
    "风险": "risk_condition",
    "综合征": "western_syndrome",
    "异常": "lab_or_imaging_finding",
}

SYNONYM_NORMALIZATION = {
    "小儿反复呼吸道感染": "儿童反复呼吸道感染",
    "小儿急性扁桃体炎": "儿童急性扁桃体炎",
    "小儿肺炎": "儿童肺炎",
    "小儿支气管肺炎": "儿童支气管肺炎",
    "过敏性鼻炎": "变应性鼻炎",
    "过敏性紫癜": "IgA血管炎",
    "紫癜性肾炎": "IgA血管炎性肾炎",
    "失眠症": "失眠障碍",
    "焦虑症": "焦虑障碍",
    "抑郁症": "抑郁障碍",
}

NON_WESTERN_HINTS = [
    "证型", "方剂", "汤", "丸", "散", "中药", "治法", "治则", "辨证",
    "病机", "加减", "处方", "用药", "剂量", "痹证",
]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def resolve_excel_path(excel_path: str | Path = DEFAULT_EXCEL) -> Path:
    path = Path(excel_path)
    if path.exists():
        return path
    desktop = Path("/Users/fangxuan/Desktop")
    for candidate in desktop.glob("**/5月5日 知识库 病名lsx.xlsx"):
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Excel knowledge base not found: {path}")


def _clean(text: Any) -> str:
    if text is None:
        return ""
    s = str(text)
    s = s.replace("**", "").replace("###", "").replace("##", "")
    s = re.sub(r"[ \t]+", " ", s)
    return s.strip()


def _split_aliases(text: str) -> list[str]:
    aliases: list[str] = []
    for part in re.split(r"[,，、/；;\n]+", text):
        p = _clean(part).strip("-：: ")
        if p and p not in aliases:
            aliases.append(p)
    return aliases


def _normalize_name_piece(piece: str) -> tuple[str, list[str]]:
    aliases: list[str] = []
    text = _clean(piece).strip("-：: ")
    text = text.strip("【】[]")
    for paren in re.findall(r"[（(]([^）)]+)[）)]", text):
        for alias in _split_aliases(paren):
            if alias and alias not in aliases:
                aliases.append(alias)
    primary = re.sub(r"[（(][^）)]*[）)]", "", text).strip()
    primary = primary.strip("【】[]")
    if not re.search(r"[\u4e00-\u9fff]", primary) and re.fullmatch(r"[A-Za-z0-9 -]{1,12}[）)]?", primary):
        if primary and primary not in aliases:
            aliases.append(primary.rstrip("）)"))
        primary = ""
    return primary, aliases


def _extract_names(cell: str) -> tuple[list[str], list[str], list[str]]:
    text = _clean(cell)
    warnings: list[str] = []
    names: list[str] = []
    aliases: list[str] = []
    if not text:
        return names, aliases, warnings

    for match in re.finditer(r"(?:主病名|病名)[：:]\s*([^\n]+)", text):
        raw = _clean(match.group(1))
        raw = re.split(r"(?:别名|基础方|方案|证型|全病程|终端药店)[：:]", raw)[0]
        for item in _split_aliases(raw):
            primary, piece_aliases = _normalize_name_piece(item)
            aliases.extend(a for a in piece_aliases if a not in aliases)
            if primary and primary not in names:
                names.append(primary)

    for match in re.finditer(r"别名[：:]\s*([^\n]+)", text):
        aliases.extend(a for a in _split_aliases(match.group(1)) if a not in aliases)

    if not names:
        first = re.split(r"\n|；|;", text)[0]
        first = re.sub(r"^(?:主病名|病名)[：:]", "", first).strip()
        for item in _split_aliases(first):
            primary, piece_aliases = _normalize_name_piece(item)
            aliases.extend(a for a in piece_aliases if a not in aliases)
            if primary and len(primary) <= 30:
                names.append(primary)

    if any(h in text for h in NON_WESTERN_HINTS):
        warnings.append("cell_contains_tcm_or_plan_noise")
    return names, aliases, warnings


def _standardize(name: str) -> tuple[str, list[str]]:
    aliases: list[str] = []
    clean = _clean(name).strip("-：: ")
    clean = re.sub(r"[（(]\s*[）)]", "", clean)
    standard = SYNONYM_NORMALIZATION.get(clean, clean)
    if standard != clean:
        aliases.append(clean)
    return standard, aliases


def _entry_type(name: str) -> str:
    for key, value in ENTRY_TYPES.items():
        if key in name:
            return value
    return "standard_western_disease"


def _allowed(name: str) -> bool:
    return not any(h in name for h in NON_WESTERN_HINTS)


def _headers(row) -> dict[str, int]:
    result = {}
    for i, cell in enumerate(row):
        value = _clean(cell.value)
        if value:
            result[value] = i
    return result


def build_disease_manifest(excel_path: str | Path = DEFAULT_EXCEL) -> dict:
    path = resolve_excel_path(excel_path)
    wb = load_workbook(path, read_only=True, data_only=True)
    diseases_by_name: dict[str, dict] = {}
    disease_occurrence_count: dict[str, int] = {}
    raw_rows = 0
    raw_cells = 0

    for ws in wb.worksheets:
        rows = ws.iter_rows()
        try:
            first = next(rows)
        except StopIteration:
            continue
        header = _headers(first)
        disease_idx = header.get("病名【诊断】")
        if disease_idx is None:
            continue
        dept_idx = header.get("科目")
        sex_idx = header.get("性别")
        age_idx = header.get("好发年龄")

        for row in rows:
            raw_rows += 1
            cell = row[disease_idx].value if disease_idx < len(row) else None
            if not cell:
                continue
            raw_cells += 1
            raw_text = _clean(cell)
            names, aliases, warnings = _extract_names(raw_text)
            for raw_name in names:
                standard, normalization_aliases = _standardize(raw_name)
                if not standard:
                    continue
                disease_occurrence_count[standard] = disease_occurrence_count.get(standard, 0) + 1
                entry = diseases_by_name.setdefault(standard, {
                    "raw_text": raw_name,
                    "standard_western_name": standard,
                    "aliases": [],
                    "department_or_category": _clean(row[dept_idx].value) if dept_idx is not None and dept_idx < len(row) else "",
                    "sex": _clean(row[sex_idx].value) if sex_idx is not None and sex_idx < len(row) else "",
                    "typical_age": _clean(row[age_idx].value) if age_idx is not None and age_idx < len(row) else "",
                    "entry_type": _entry_type(standard),
                    "m1_diagnosis_entry_allowed": _allowed(standard),
                    "needs_external_enrichment": True,
                    "cleaning_warnings": [],
                })
                for alias in aliases + normalization_aliases:
                    if alias and alias != standard and alias not in entry["aliases"]:
                        entry["aliases"].append(alias)
                for warning in warnings:
                    if warning not in entry["cleaning_warnings"]:
                        entry["cleaning_warnings"].append(warning)
                if not _allowed(standard):
                    entry["m1_diagnosis_entry_allowed"] = False
                    if "non_western_or_plan_noise" not in entry["cleaning_warnings"]:
                        entry["cleaning_warnings"].append("non_western_or_plan_noise")

    diseases = sorted(diseases_by_name.values(), key=lambda x: x["standard_western_name"])
    manifest = {
        "created_at": _now(),
        "source_file": str(path),
        "raw_row_count": raw_rows,
        "raw_disease_cell_count": raw_cells,
        "cleaned_disease_count": len(diseases),
        "duplicate_disease_count": sum(max(0, count - 1) for count in disease_occurrence_count.values()),
        "non_western_or_blocked_count": sum(1 for item in diseases if not item.get("m1_diagnosis_entry_allowed", True)),
        "diseases": diseases,
    }
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    result = build_disease_manifest()
    print(json.dumps({
        "raw_disease_cell_count": result["raw_disease_cell_count"],
        "cleaned_disease_count": result["cleaned_disease_count"],
    }, ensure_ascii=False, indent=2))
